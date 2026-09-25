"""Slide-aligned lecture notes: a screen-capture recording + its transcript -> a
Word document with one entry per slide (the 720p slide image, its timecode
range, and the transcript spoken while that slide was up).

Slides are found by frame-differencing the video (robust for text slides, where
scene-score detectors fail), then near-identical consecutive frames are merged
with a perceptual hash so a moving laser-pointer dot doesn't split one slide in
two -- while genuine animation builds are kept as separate steps.

  uv run python slidedoc.py \
      --video lecture.mp4 --transcript out/<slug>.segments.json \
      --out out/<slug>.slides.docx \
      --title "..." --lecturer "..." --date "18 September 2026" --course "BIOT 505"

Everything it reads and writes is transient/private; nothing is committed.
"""

import argparse
import json
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path

import numpy as np
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def hms(t):
    t = int(round(t))
    return f"{t // 3600}:{(t % 3600) // 60:02d}:{t % 60:02d}"


def video_duration(video):
    out = subprocess.run(
        ["ffprobe", "-v", "error", "-show_entries", "format=duration",
         "-of", "default=noprint_wrappers=1:nokey=1", video],
        capture_output=True, text=True, check=True)
    return float(out.stdout.strip())


def detect_changes(video, work, step=2, w=96, h=54, min_gap=8.0):
    """Timestamps (s) where the slide image changes, by mean-abs-diff of dense
    downscaled grayscale frames."""
    raw = work / "frames.raw"
    subprocess.run(
        ["ffmpeg", "-y", "-hide_banner", "-loglevel", "error", "-i", video,
         "-vf", f"fps=1/{step},scale={w}:{h},format=gray",
         "-f", "rawvideo", "-pix_fmt", "gray", str(raw)], check=True)
    a = np.fromfile(raw, dtype=np.uint8)
    n = a.size // (w * h)
    a = a[:n * w * h].reshape(n, w * h).astype(np.int16)
    mad = np.abs(np.diff(a, axis=0)).mean(axis=1)
    # threshold: well above the static baseline, below real slide-flip spikes
    med, q75 = np.percentile(mad, 50), np.percentile(mad, 75)
    thr = max(6.0, med + 6 * (q75 - med + 1e-6))
    changes = []
    for i in np.where(mad > thr)[0]:
        t = float((i + 1) * step)
        if changes and t - changes[-1] < min_gap:
            continue
        changes.append(t)
    return changes, float(thr)


def extract_frames(video, bounds, slides_dir):
    """One representative full-res frame per [start,end) segment, sampled late
    enough to catch the fully-built slide."""
    slides_dir.mkdir(parents=True, exist_ok=True)
    segs = []
    for i in range(len(bounds) - 1):
        s, e = bounds[i], bounds[i + 1]
        dur = e - s
        t = min(max(s + 1.5, s + dur * 0.7), e - 1.5) if dur > 3 else s + dur / 2
        out = slides_dir / f"slide_{i + 1:03d}.jpg"
        subprocess.run(
            ["ffmpeg", "-y", "-hide_banner", "-loglevel", "error", "-ss", f"{t:.2f}",
             "-i", video, "-frames:v", "1", "-q:v", "3", str(out)], check=True)
        segs.append({"start": s, "end": e, "frame_t": t, "img": str(out)})
    return segs


def dhash_bits(path, hx=16, hy=16):
    im = Image.open(path).convert("L").resize((hx + 1, hy), Image.LANCZOS)
    px = np.asarray(im, dtype=np.int16)
    return (px[:, :-1] < px[:, 1:]).flatten()


def dedup(segs, thr=12):
    """Merge consecutive near-identical slides (pointer-only changes)."""
    out = []
    last_h = None
    for s in segs:
        h = dhash_bits(s["img"])
        if out and last_h is not None and int(np.count_nonzero(h != last_h)) <= thr:
            out[-1]["end"] = s["end"]          # extend, keep first image
        else:
            out.append(dict(s))
            last_h = h
    return out


def align(segs, cues):
    """Attach each transcript cue to the slide whose interval holds its midpoint."""
    for s in segs:
        s["cues"] = []
    for c in cues:
        mid = (c["start"] + c["end"]) / 2
        placed = False
        for s in segs:
            if s["start"] <= mid < s["end"]:
                s["cues"].append(c)
                placed = True
                break
        if not placed and segs:            # tail cue -> last slide
            segs[-1]["cues"].append(c)
    return segs


def clean_image(path, max_w=1280):
    """Re-encode to a clean baseline JPEG python-docx will accept (ffmpeg's
    JPEGs occasionally trip its strict header parser)."""
    im = Image.open(path).convert("RGB")
    if im.width > max_w:
        im = im.resize((max_w, round(im.height * max_w / im.width)), Image.LANCZOS)
    buf = BytesIO()
    im.save(buf, format="JPEG", quality=85)
    buf.seek(0)
    return buf


def build_docx(segs, meta, out_path):
    doc = Document()
    normal = doc.styles["Normal"].font
    normal.name = "Calibri"
    normal.size = Pt(11)

    title = doc.add_heading(meta["title"], level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in [meta.get("lecturer", ""), meta.get("course", ""), meta.get("date", "")]:
        if not line:
            continue
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.size = Pt(12)

    n_flag = sum(1 for s in segs for c in s["cues"] if c.get("needs_review"))
    note = doc.add_paragraph()
    r = note.add_run(
        f"{len(segs)} slides · transcript by gpt-transcribe (cross-checked) · "
        f"slide timings from the recording · {n_flag} passage(s) flagged for review "
        f"are shown in italics. Personal study notes.")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for i, s in enumerate(segs, 1):
        if i > 1:
            doc.add_page_break()
        head = doc.add_heading(level=1)
        rh = head.add_run(f"Slide {i}")
        rt = head.add_run(f"    {hms(s['start'])} – {hms(s['end'])}")
        rt.font.size = Pt(12)
        rt.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
        doc.add_picture(clean_image(s["img"]), width=Inches(6.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

        if not s["cues"]:
            e = doc.add_paragraph("(no speech recorded on this slide)")
            e.runs[0].font.italic = True
            e.runs[0].font.color.rgb = RGBColor(0x99, 0x99, 0x99)
            continue
        para = doc.add_paragraph()
        for c in s["cues"]:
            run = para.add_run(c["text"].strip() + " ")
            if c.get("needs_review"):
                run.font.italic = True
                run.font.color.rgb = RGBColor(0x9A, 0x63, 0x00)

    doc.save(out_path)
    return out_path, n_flag


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--video", required=True)
    ap.add_argument("--transcript", required=True, help="<slug>.segments.json from run.py")
    ap.add_argument("--out", required=True)
    ap.add_argument("--title", default="Lecture")
    ap.add_argument("--lecturer", default="")
    ap.add_argument("--course", default="")
    ap.add_argument("--date", default="")
    ap.add_argument("--step", type=float, default=2.0, help="sampling seconds for change detection")
    ap.add_argument("--workdir", default="")
    args = ap.parse_args()

    video = str(Path(args.video).expanduser())
    cues = json.load(open(Path(args.transcript).expanduser()))
    dur = video_duration(video)

    tmp = Path(args.workdir).expanduser() if args.workdir else Path(tempfile.mkdtemp(prefix="slidedoc_"))
    tmp.mkdir(parents=True, exist_ok=True)
    slides_dir = tmp / "slides"

    print(f"[1/5] Detecting slide changes ({dur/60:.1f} min video)…")
    changes, thr = detect_changes(video, tmp, step=args.step)
    bounds = [0.0] + changes + [dur]
    print(f"      {len(changes)} changes (threshold {thr:.1f}) -> {len(bounds)-1} raw segments")

    print("[2/5] Extracting full-res slide frames…")
    segs = extract_frames(video, bounds, slides_dir)

    print("[3/5] De-duplicating pointer-only repeats…")
    segs = dedup(segs)
    print(f"      -> {len(segs)} slides")

    print(f"[4/5] Aligning {len(cues)} transcript cues to slides…")
    segs = align(segs, cues)

    print("[5/5] Writing Word document…")
    meta = {"title": args.title, "lecturer": args.lecturer,
            "course": args.course, "date": args.date}
    out_path = str(Path(args.out).expanduser())
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    path, n_flag = build_docx(segs, meta, out_path)
    print(f"\nWrote {path}  ({len(segs)} slides, {n_flag} flagged passages)")


if __name__ == "__main__":
    main()
